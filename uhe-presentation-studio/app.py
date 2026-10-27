#!/usr/bin/env python3
"""UHE Presentation Studio: local-only project intake, evidence and PPTX/PDF builder."""
from __future__ import annotations
import csv, hashlib, json, mimetypes, os, re, shutil, subprocess, sys, tempfile, threading, uuid
from datetime import datetime
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlparse, parse_qs, unquote

ROOT=Path(__file__).resolve().parent
DATA=Path(os.environ.get('UHE_STUDIO_DATA',ROOT/'data')).resolve()
PROJECTS=DATA/'projects'; PROFILES=DATA/'profiles'; WEB=ROOT/'web'
for p in (PROJECTS,PROFILES): p.mkdir(parents=True,exist_ok=True)
ALLOWED={'.pdf','.docx','.pptx','.ppt','.xlsx','.csv','.txt','.md','.png','.jpg','.jpeg','.webp','.tif','.tiff','.bmp','.svg','.gltf','.glb','.obj','.stl','.fbx','.blend'}
MAX_UPLOAD=200*1024*1024
DEFAULT_PROFILE={'name':'UHE Default','page_size':'16:9','width_in':13.333,'height_in':7.5,'language':'en','rtl':False,'colors':{'ink':'152238','accent':'D7653B','secondary':'1B7B78','paper':'F6F2EA','white':'FFFFFF','muted':'677386'},'font_title':'Aptos Display','font_body':'Aptos','title_pt':30,'body_pt':17,'margin_in':0.55,'footer':'UNIQUE HASHTAGGER EVENTS • DRAFT','page_numbers':True,'animation':'none','image_fit':'contain'}
if not (PROFILES/'uhe-default.json').exists(): (PROFILES/'uhe-default.json').write_text(json.dumps(DEFAULT_PROFILE,indent=2),'utf-8')

def slug(s):
    s=re.sub(r'[^\w.-]+','-',str(s),flags=re.UNICODE).strip('-._').lower()[:70]
    return s or 'project-'+uuid.uuid4().hex[:8]

def safe_project(pid):
    p=(PROJECTS/slug(pid)).resolve()
    if PROJECTS not in p.parents: raise ValueError('Invalid project id')
    return p

def read_json(p,default):
    try: return json.loads(p.read_text('utf-8'))
    except Exception: return default

def write_json(p,obj): p.parent.mkdir(parents=True,exist_ok=True); p.write_text(json.dumps(obj,ensure_ascii=False,indent=2),'utf-8')

def archive_slides_if_changed(project,deck):
    previous=read_json(project/'slides.json',None)
    if previous and previous!=deck:
        stamp=datetime.now().strftime('%Y%m%d-%H%M%S-%f');archive=project/'revisions'/stamp;archive.mkdir(parents=True,exist_ok=True);write_json(archive/'slides.json',previous)

def extract_file(path, project_root=None):
    ext=path.suffix.lower(); text=''; refs=[]; issue=None; media=[]
    try:
        if ext=='.pdf':
            import fitz
            d=fitz.open(path)
            ocr_used=False
            pdf_text=' '.join(pg.get_text('text') for pg in d[:50])
            letter_count=sum(ch.isalpha() for ch in pdf_text)
            force_ocr=len(pdf_text)>500 and pdf_text.count(' ')/max(letter_count,1)>.45
            seen_image_xrefs=set()
            for i,page in enumerate(d[:50]):
                t=page.get_text('text').strip()
                quality_lines=[line for line in t.splitlines() if not re.fullmatch(r'[\sA-Z0-9/&.\-–—]+',line)]
                spaced=len(re.findall(r'\b[A-Za-z](?:\s[A-Za-z]){3,}\b','\n'.join(quality_lines)))
                if len(t)>150 and (force_ocr or spaced>=1) and shutil.which('tesseract'):
                    try:
                        with tempfile.TemporaryDirectory() as td:
                            image=Path(td)/'page.png';page.get_pixmap(matrix=fitz.Matrix(3,3),alpha=False).save(image)
                            ocr=subprocess.run([shutil.which('tesseract'),str(image),'stdout','-l','eng'],capture_output=True,text=True,timeout=60)
                            if ocr.returncode==0 and len(ocr.stdout.strip())>40:
                                t=ocr.stdout.strip();ocr_used=True
                                refs.append({'where':f'page {i+1} (OCR; verify visually)','text':t})
                                continue
                    except Exception: pass
                if t: refs.append({'where':f'page {i+1}','text':t})
                for j,img in enumerate(page.get_image_info(xrefs=True)):
                    if len(media)>=12: break
                    try:
                        xref=img.get('xref',0)
                        if not xref: continue
                        if xref in seen_image_xrefs: continue
                        seen_image_xrefs.add(xref)
                        pix=fitz.Pixmap(d,xref)
                        if pix.width<800 or pix.height<500: continue
                        pix=fitz.Pixmap(fitz.csRGB,pix) if pix.n>4 else pix
                        assetroot=(Path(project_root)/'derived_visuals') if project_root else path.parent/'assets'
                        out=assetroot/f'{path.stem}-p{i+1}-img{j+1}.png'; out.parent.mkdir(parents=True,exist_ok=True)
                        pix.save(out); media.append(str(out.relative_to(project_root)) if project_root else str(out.relative_to(path.parent)))
                    except Exception: pass
            if len(d)>50: issue='File exceeds the 50-page source limit; only the first 50 pages were extracted.'; refs=refs[:50]
        elif ext=='.docx':
            from docx import Document
            d=Document(path); body=[p.text.strip() for p in d.paragraphs if p.text.strip()]
            for n,t in enumerate(body): refs.append({'where':f'paragraph {n+1}','text':t})
            for ti,table in enumerate(d.tables,1):
                for ri,row in enumerate(table.rows,1):
                    t=' | '.join(c.text.strip().replace('\n',' / ') for c in row.cells)
                    if t.strip(' |'): refs.append({'where':f'table {ti}, row {ri}','text':t})
        elif ext=='.pptx':
            from pptx import Presentation
            d=Presentation(path)
            # python-pptx slide collections do not support slice indexing; list first.
            for i,slide in enumerate(list(d.slides)[:50],1):
                bits=[]
                for sh in slide.shapes:
                    if getattr(sh,'has_text_frame',False) and sh.text.strip(): bits.append(sh.text.strip())
                    if getattr(sh,'has_table',False):
                        for row in sh.table.rows: bits.append(' | '.join(c.text.strip() for c in row.cells))
                    if sh.shape_type==13:
                        try:
                            assetroot=(Path(project_root)/'derived_visuals') if project_root else path.parent/'assets'
                            out=assetroot/f'{path.stem}-s{i}-img{len(media)+1}.png';out.parent.mkdir(parents=True,exist_ok=True);out.write_bytes(sh.image.blob);media.append(str(out.relative_to(project_root)) if project_root else str(out.relative_to(path.parent)))
                        except Exception: pass
                if bits: refs.append({'where':f'slide {i}','text':'\n'.join(bits)})
            if len(d.slides)>50: issue='File exceeds the 50-slide source limit; only the first 50 slides were extracted.'
        elif ext=='.ppt':
            soffice=shutil.which('soffice')
            if not soffice: issue='Legacy PPT needs LibreOffice for extraction.'
            else:
                with tempfile.TemporaryDirectory() as td:
                    subprocess.run([soffice,'--headless','--convert-to','pptx','--outdir',td,str(path)],capture_output=True,timeout=90)
                    converted=Path(td)/(path.stem+'.pptx')
                    if converted.exists():
                        from pptx import Presentation
                        d=Presentation(converted)
                        for i,slide in enumerate(d.slides,1):
                            t='\n'.join(sh.text.strip() for sh in slide.shapes if getattr(sh,'has_text_frame',False) and sh.text.strip())
                            if t: refs.append({'where':f'slide {i}','text':t})
                    else: issue='LibreOffice could not read this legacy PPT.'
        elif ext=='.xlsx':
            from openpyxl import load_workbook
            w=load_workbook(path,read_only=True,data_only=True)
            for sh in w.worksheets:
                for row in sh.iter_rows():
                    vals=[str(c.value) for c in row if c.value is not None]
                    if vals: refs.append({'where':f'{sh.title}!{row[0].row}','text':' | '.join(vals)})
        elif ext=='.csv':
            with path.open(encoding='utf-8-sig',errors='replace',newline='') as f:
                for i,row in enumerate(csv.reader(f),1): refs.append({'where':f'row {i}','text':' | '.join(row)})
        elif ext in ('.txt','.md'):
            for i,line in enumerate(path.read_text('utf-8',errors='replace').splitlines(),1):
                if line.strip(): refs.append({'where':f'line {i}','text':line.strip()})
        elif ext=='.svg':
            import xml.etree.ElementTree as ET
            root=ET.parse(path).getroot(); labels=[x.text.strip() for x in root.iter() if x.tag.rsplit('}',1)[-1] in {'text','title','desc'} and x.text and x.text.strip()]
            if labels: refs.append({'where':'SVG text labels','text':'\n'.join(labels)})
            media.append(str(path.relative_to(project_root)) if project_root else path.name)
        elif ext in {'.gltf','.glb','.obj','.stl','.fbx','.blend'}:
            issue='3D source inventoried and preserved; this build does not render or extract this 3D format.'
            if ext=='.gltf':
                try:
                    g=json.loads(path.read_text('utf-8'));refs.append({'where':'glTF metadata','text':f"Meshes: {len(g.get('meshes',[]))}; nodes: {len(g.get('nodes',[]))}; scenes: {len(g.get('scenes',[]))}"})
                except Exception: issue='Could not read glTF JSON; 3D source preserved.'
        elif ext in {'.png','.jpg','.jpeg','.webp','.tif','.tiff','.bmp'}:
            media.append(str(path.relative_to(project_root)) if project_root else path.name)
            exe=shutil.which('tesseract')
            if exe:
                r=subprocess.run([exe,str(path),'stdout','-l','eng'],capture_output=True,text=True,timeout=45)
                if r.returncode==0 and r.stdout.strip(): refs.append({'where':'image OCR','text':r.stdout.strip()})
            else: issue='Image inventoried; OCR is unavailable because Tesseract is not installed.'
        if ext=='.pdf' and refs:
            sample=' '.join(x['text'] for x in refs)
            if ocr_used: issue='OCR fallback was applied to split PDF text; verify every extracted claim visually against the source page.'
            elif len(re.findall(r'\b[A-Za-z](?:\s[A-Za-z]){3,}\b','\n'.join(line for line in sample.splitlines() if not re.fullmatch(r'[\sA-Z0-9/&.\-–—]+',line))))>=2 or (len(sample)>500 and sample.count(' ')/max(sum(ch.isalpha() for ch in sample),1)>.45):
                issue='PDF text extraction appears degraded (split or corrupted words); verify the source visually before using its claims.'
    except Exception as e: issue=f'Could not read file: {type(e).__name__}: {e}'
    return {'text':'\n'.join(x['text'] for x in refs),'references':refs,'media':media,'issue':issue}

def inventory_project(p):
    files=[]
    for f in sorted(p.rglob('*')):
        if not f.is_file() or f.name in {'inventory.json','slides.json'} or 'output' in f.parts or 'derived_visuals' in f.parts: continue
        rel=str(f.relative_to(p)); digest=hashlib.sha256(f.read_bytes()).hexdigest()
        item={'path':rel,'name':f.name,'ext':f.suffix.lower(),'bytes':f.stat().st_size,'sha256':digest,'modified':datetime.fromtimestamp(f.stat().st_mtime).isoformat(timespec='seconds')}
        if f.suffix.lower() in ALLOWED:
            e=extract_file(f,p)
            item.update(e)
            item['kind']=('visual' if f.suffix.lower() in {'.png','.jpg','.jpeg','.webp','.tif','.tiff','.bmp','.svg'} else '3d' if f.suffix.lower() in {'.gltf','.glb','.obj','.stl','.fbx','.blend'} else 'document')
        else: item.update({'text':'','references':[],'media':[],'issue':'Unsupported file type; preserved in the project folder.','kind':'other'})
        files.append(item)
    seen={}; duplicates=[]
    for x in files:
        if x['sha256'] in seen: duplicates.append({'files':[seen[x['sha256']],x['path']],'type':'exact duplicate'})
        else: seen[x['sha256']]=x['path']
    for f in files:
        if f['kind']=='visual':
            try:
                from PIL import Image
                with Image.open(p/f['path']) as im:
                    f['dimensions']=[im.width,im.height]
                    if im.width<900 or im.height<600: f['issue']='Low-resolution image ('+str(im.width)+'×'+str(im.height)+' px); inspect before using in a full-slide visual.'
            except Exception: pass
    media_assets=[]
    for f in files:
        for m in f.get('media',[]):
            match=re.search(r'-(?:s|p)(\d+)-img',Path(m).name)
            where=(f"slide {match.group(1)} image" if match and '-s' in Path(m).name else f"page {match.group(1)} image" if match else 'image')
            asset={'path':m,'source':f"{f['path']} · {where}",'type':Path(m).suffix.lower(),'bytes':(p/m).stat().st_size if (p/m).exists() else 0,'unverified':True}
            try:
                from PIL import Image
                with Image.open(p/m) as im:
                    asset['dimensions']=[im.width,im.height]
                    if im.width<1200 or im.height<700: asset['issue']=f'Low resolution: {im.width}×{im.height}px'
            except Exception: pass
            media_assets.append(asset)
    return {'project_id':p.name,'name':p.name.replace('-',' ').title(),'scanned_at':datetime.now().isoformat(timespec='seconds'),'files':files,'duplicates':duplicates,'media_assets':media_assets}

def facts_for(inv):
    facts=[];seen=set()
    for f in inv.get('files',[]):
        if (f.get('issue') or '').startswith('PDF text extraction appears degraded'): continue
        for ref in f.get('references',[]):
            txt=ref['text']
            # Sentence-level sourcing keeps every generated claim anchored to a source location.
            for sentence in re.split(r'(?<=[.!?؟])\s+|\n+',txt):
                sentence=re.sub(r'\s+',' ',sentence).strip(' •\t')
                prefix=re.sub(r'[^A-Za-z]','',sentence[:50]).upper()
                normalized=re.sub(r'[^A-Z0-9]','',sentence.upper())
                if prefix.startswith(('INTERNALDRAFT','CONFLICT','BLOCKER','AMBIGUITY','SOURCE','DRAWING','PERSPECTIVE','CONFIRM','ASSUMPTIONS','EDITABLEMODEL','ALLRATESINTENTIONALLYBLANK')): continue
                if any(flag in normalized for flag in ('INTERNALDRAFT','INTERNALMANIFEST','SHA256','CONFLICT','BLOCKER','AMBIGUITY','EDITABLESOURCE','MODELSPACEGLB','REVIEWISSUE','NOTSURVEYED','APPROVALSPENDING','ALLRATESINTENTIONALLYBLANK','SURVEYREQUIRED','NEXTACTION','RESOLUTIONOFCONTRADICTORY','RESOLUTIONOFEQUIPMENT','OBTAINING6BRANCHSURVEYS','SUBSTANTIATIONOFSUPPLIERRATES','RECEIPTOFOFFICIAL','REMAINUNVERIFIED','STILLREQUIRED')): continue
                if normalized in {'REQUIREMENTCONTROL','VERIFIEDSTATUSRELEASEBLOCKER','PROPOSALDEADLINE','SIXBRANCHSCOPE','SITEGEOMETRY','EQUIPMENTQUANTITIES','OPERATINGPERIOD','BRANDGUIDELINES','COMMERCIALRELEASE','LIGHTINGPERFORMANCEMATERIALCOMPLIANCE','FACADEWASHGRAPHICSSUBSTRATE','SAFETYPOWERWEATHERRATING','DESIGNDELIVERABLESANDAPPROVALS','PHASEDELIVERABLEPACKPERBRANCHCBUAEHOLDPOINT'}: continue
                letters=''.join(ch for ch in sentence if ch.isalpha())
                if letters and letters.isupper() and not any(ch.isdigit() for ch in sentence) and len(sentence)<80: continue
                if re.search(r'(?i)\b(?:supplier\s+buying\s+cost|buying\s+cost|margin|markup|negotiation\s+strategy|internal\s+cost)\b',sentence): continue
                key=re.sub(r'\W+','',sentence.lower())
                if key in seen: continue
                seen.add(key)
                if len(sentence)>18: facts.append({'text':sentence,'source':f"{f['path']} · {ref['where']}"})
    return facts

def make_draft(inv,opts):
    name=opts.get('project_name') or inv.get('name','Project'); lang=opts.get('language','en'); kind=opts.get('proposal_type','combined')
    facts=facts_for(inv); cap=int(opts.get('max_slides',12)); cap=max(4,min(50,cap))
    slides=[{'layout':'cover','title':name,'body':[f"{kind.title()} proposal",'Draft • source-linked for review'],'sources':[],'media':[]}]
    selected_groups={'Brief & intent':[],'Design direction':[],'Technical scope':[],'Delivery & operations':[]}
    delivery=re.compile(r'\b(install|dismant|commission|operate|operation|schedule|mobilis|handover|maintenance|opening|closing|phase|day\s*\d)\b',re.I)
    technical=re.compile(r'\b(kva|kw|ip\d{2}|cable|power|electrical|fixture|equipment|capacity|accessib|clearance|fire|safety|ramp|load|dimension|meter|sqm|m²|seat|booth|screen|lighting|material|specification|structural|route|circulation|toilet|wash station|console|rack|place|proxy)\b',re.I)
    creative=re.compile(r'\b(concept|design|theme|palette|colour|color|hospitality|guest journey|visitor|experience|atmosphere|story|identity|visual|lighting|material|gathering|majlis)\b',re.I)
    for fact in facts:
        txt=fact['text']
        group='Delivery & operations' if delivery.search(txt) else 'Technical scope' if technical.search(txt) else 'Design direction' if creative.search(txt) else 'Brief & intent'
        if len(selected_groups[group])<5: selected_groups[group].append(fact)
    groups=[('Brief & intent',selected_groups['Brief & intent'],'statement'),('Design direction',selected_groups['Design direction'],'evidence'),('Technical scope',selected_groups['Technical scope'],'scope'),('Delivery & operations',selected_groups['Delivery & operations'],'sequence')]
    if kind=='design': groups=[x for x in groups if x[0] in ('Brief & intent','Design direction')]
    elif kind=='technical': groups=[x for x in groups if x[0] in ('Brief & intent','Technical scope','Delivery & operations')]
    for title,items,layout in groups:
        if len(slides)>=cap-2: break
        if items:
            slides.append({'layout':layout,'title':title,'body':[x['text'] for x in items[:5]],'sources':[x['source'] for x in items[:5]],'media':[]})
    selected=set(opts.get('selected_media',[]))
    images=[x for x in inv.get('media_assets',[]) if x.get('path') in selected]
    visual_slots=max(0,min(5,cap-len(slides)-2))
    for i,x in enumerate(images[:visual_slots],1):
        slides.append({'layout':'visual','title':f'Design evidence · {i:02d}','body':['Source view reproduced from the cited project presentation.'],'sources':[x['source']],'media':[x['path']]})
    unresolved=[]
    if not facts: unresolved.append('No machine-readable proposal text found. Add an accessible source or extract text manually.')
    if not images: unresolved.append('No project visuals were approved for this draft. Preview and select authentic, approved photography, renders, plans or brand assets before using them.')
    for f in inv.get('files',[]):
        if (f.get('issue') or '').startswith('OCR fallback'):
            unresolved.append(f"Verify OCR-extracted text from {f['path']} against the original page images before client use.")
        raw=' '.join(ref['text'] for ref in f.get('references',[]));norm=re.sub(r'[^A-Z0-9]','',raw.upper())
        if any(flag in norm for flag in ('CONFLICT','BLOCKER','AMBIGUITY','NOTSURVEYED','APPROVALSPENDING','SURVEYREQUIRED','ALLRATESINTENTIONALLYBLANK','REVIEWISSUE')):
            unresolved.append(f"{f['path']} contains unresolved site, revision, survey, approval or commercial-release status notes; resolve them before client use.")
        for ref in f.get('references',[]):
            for label in ('CONFLICT:','BLOCKER:','AMBIGUITY:'):
                if label in ref['text'].upper(): unresolved.append(f"Source flag in {f['path']} · {ref['where']}: review the marked {label[:-1].lower()} items.");break
    # Do not infer conflicts from differing dates, prices or quantities without context; explicit source conflict markers are surfaced above.
    slides.append({'layout':'review','title':'Proposal basis','body':['Confirm the controlling source revisions, final scope, site inputs, approvals and delivery dependencies with the client before issue.'],'sources':[],'media':[]})
    slides.append({'layout':'close','title':'Next step','body':['Confirm the open items, approve the selected source revisions, then issue the reviewed proposal.'],'sources':[],'media':[]})
    return {'project_name':name,'proposal_type':kind,'language':lang,'profile_name':opts.get('profile_name','UHE Default'),'visual_style':opts.get('visual_style','layered'),'status':'DRAFT • NOT APPROVED','slides':slides,'review_items':unresolved,'generated_at':datetime.now().isoformat(timespec='seconds')}

def build_pptx(p,deck,profile):
    from pptx import Presentation
    from pptx.util import Inches,Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
    from pptx.enum.shapes import MSO_SHAPE
    r=Presentation(); r.slide_width=Inches(float(profile.get('width_in',13.333)));r.slide_height=Inches(float(profile.get('height_in',7.5)))
    c=dict(profile['colors'])
    if profile.get('name','UHE Default')=='UHE Default' and deck.get('visual_style')=='quiet': c.update({'ink':'17324D','accent':'C39A52','secondary':'40677A','paper':'F7F6F2','muted':'687685'})
    if profile.get('name','UHE Default')=='UHE Default' and deck.get('visual_style')=='bold': c.update({'ink':'141C2E','accent':'F0512A','secondary':'3458E5','paper':'F7F8FA','muted':'5C6678'})
    W=float(profile.get('width_in',13.333));H=float(profile.get('height_in',7.5)); m=float(profile.get('margin_in',.55))
    def rgb(h): return RGBColor.from_string(h.replace('#','').upper())
    def source_name(source):
        filename=Path(source.split(' · ')[0]).name
        aliases={'CBUAE_BPM000921_—_Design_Proposal_Draft.pdf':'CBUAE Design Proposal Draft','MOHESR_Design_Proposal_SIGNED_STAMPED.pdf':'MOHESR Design Proposal','MOHESR_Technical_Proposal_SIGNED_STAMPED.pdf':'MOHESR Technical Proposal','MOHESR_Master_Plan_R01(2).pdf':'MOHESR Master Plan R01','NOVA70_CURRENT_OPTION_F_DESIGN_ATLAS.pptx':'NOVA70 Design Atlas','NOVA71_FOH_VIP_COORDINATED_DESIGN_SUPPLEMENT.pptx':'NOVA71 FOH/VIP Supplement'}
        if filename in aliases: return aliases[filename]
        return re.sub(r'[_%]+',' ',Path(filename).stem).replace('SIGNED STAMPED','').strip()
    blank=r.slide_layouts[6]
    for i,s in enumerate(deck['slides'],1):
        unique_sources=list(dict.fromkeys(s.get('sources',[])))
        source_numbers={source:n for n,source in enumerate(unique_sources,1)}
        def source_marker(j):
            sources=s.get('sources',[])
            return f" [{source_numbers[sources[j]]}]" if j<len(sources) and sources[j] in source_numbers else ''
        sl=r.slides.add_slide(blank); bg=sl.background.fill;bg.solid();bg.fore_color.rgb=rgb(c['ink'] if s['layout']=='statement' else c['paper'])
        # Keep the tutorial's layered depth on the cover and use distinct, restrained structures inside the story.
        if s['layout'] in ('cover','close'):
            sh=sl.shapes.add_shape(MSO_SHAPE.OVAL, Inches(W-3.7), Inches(-.7), Inches(4.2), Inches(4.2));sh.fill.solid();sh.fill.fore_color.rgb=rgb(c['secondary']);sh.line.fill.background()
            sh=sl.shapes.add_shape(MSO_SHAPE.OVAL, Inches(W-2.65), Inches(.35), Inches(3.5), Inches(3.5));sh.fill.solid();sh.fill.fore_color.rgb=rgb(c['accent']);sh.fill.transparency=20;sh.line.fill.background()
            band=sl.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,Inches(H-.22),Inches(W),Inches(.22));band.fill.solid();band.fill.fore_color.rgb=rgb(c['accent']);band.line.fill.background()
        elif s['layout']=='evidence':
            rail=sl.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,Inches(.16),Inches(H));rail.fill.solid();rail.fill.fore_color.rgb=rgb(c['accent']);rail.line.fill.background()
            orb=sl.shapes.add_shape(MSO_SHAPE.OVAL, Inches(W-1.25), Inches(-.28), Inches(1.5), Inches(1.5));orb.fill.solid();orb.fill.fore_color.rgb=rgb(c['secondary']);orb.line.fill.background()
        elif s['layout']=='statement':
            band=sl.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,Inches(.18),Inches(H));band.fill.solid();band.fill.fore_color.rgb=rgb(c['accent']);band.line.fill.background()
            marker=sl.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(W-.9),Inches(.65),Inches(.34),Inches(.08));marker.fill.solid();marker.fill.fore_color.rgb=rgb(c['secondary']);marker.line.fill.background()
        elif s['layout']=='scope':
            bar=sl.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,Inches(W),Inches(.12));bar.fill.solid();bar.fill.fore_color.rgb=rgb(c['accent']);bar.line.fill.background()
        elif s['layout']=='sequence':
            rail=sl.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(m+.17),Inches(2.42),Inches(.035),Inches(3.75));rail.fill.solid();rail.fill.fore_color.rgb=rgb(c['accent']);rail.line.fill.background()
        x=Inches(m); y=Inches(1.05 if s['layout'] not in ('cover','close') else 1.65); width=Inches(W-2*m)
        title=sl.shapes.add_textbox(x,y,width,Inches(1.15));tf=title.text_frame;tf.clear();tf.word_wrap=True
        p0=tf.paragraphs[0];p0.text=s['title'];p0.font.name=profile['font_title'];p0.font.size=Pt(profile['title_pt'] if s['layout'] in ('cover','close') else profile['title_pt']*.78);p0.font.bold=True;p0.font.color.rgb=rgb(c['paper'] if s['layout']=='statement' else c['ink'])
        if profile.get('rtl') or deck.get('language')=='ar': p0._p.get_or_add_pPr().set('rtl','1');p0.alignment=PP_ALIGN.RIGHT
        lines=s.get('body',[]);top=2.28;bodycolor=c['paper'] if s['layout']=='statement' else c['muted'] if s['layout']=='review' else c['ink']
        if s['layout']=='scope':
            for j,line in enumerate(lines):
                col=j%2;row=j//2;bx=x+Inches(col*(W/2));by=Inches(top+row*.95);bw=Inches(W/2-m-.12)
                bullet=sl.shapes.add_shape(MSO_SHAPE.OVAL,bx,by+Inches(.08),Inches(.12),Inches(.12));bullet.fill.solid();bullet.fill.fore_color.rgb=rgb(c['accent'] if j%2==0 else c['secondary']);bullet.line.fill.background()
                box=sl.shapes.add_textbox(bx+Inches(.24),by,bw-Inches(.24),Inches(.82));tf=box.text_frame;tf.clear();tf.word_wrap=True;par=tf.paragraphs[0];par.text=str(line)+source_marker(j);par.font.name=profile['font_body'];par.font.size=Pt(profile['body_pt']-1);par.font.color.rgb=rgb(bodycolor);par.space_after=Pt(6)
        elif s['layout']=='sequence':
            rowh=3.65/max(1,len(lines))
            for j,line in enumerate(lines):
                by=Inches(top+.18+j*rowh);dot=sl.shapes.add_shape(MSO_SHAPE.OVAL,x,by,Inches(.42),Inches(.42));dot.fill.solid();dot.fill.fore_color.rgb=rgb(c['accent']);dot.line.fill.background();tb=dot.text_frame;tb.text=f'{j+1:02d}';tb.paragraphs[0].font.size=Pt(9);tb.paragraphs[0].font.bold=True;tb.paragraphs[0].font.color.rgb=rgb(c['white']);tb.paragraphs[0].alignment=PP_ALIGN.CENTER
                box=sl.shapes.add_textbox(x+Inches(.72),by-Inches(.01),Inches(W-2*m-.7),Inches(rowh-.05));tf=box.text_frame;tf.clear();tf.word_wrap=True;par=tf.paragraphs[0];par.text=str(line)+source_marker(j);par.font.name=profile['font_body'];par.font.size=Pt(profile['body_pt']-1);par.font.color.rgb=rgb(bodycolor);par.vertical_anchor=MSO_ANCHOR.MIDDLE
        else:
            bodywidth=Inches(W*.34 if s['layout']=='visual' else W-2*m)
            box=sl.shapes.add_textbox(x, Inches(3.0 if s['layout'] in ('cover','close') else top), bodywidth, Inches(H-3.25));tf=box.text_frame;tf.clear();tf.word_wrap=True;tf.margin_left=Inches(.04);tf.margin_right=Inches(.04)
            for j,line in enumerate(lines):
                marker=source_marker(j)
                par=tf.paragraphs[0] if j==0 else tf.add_paragraph();par.text=(('• '+str(line)+marker) if s['layout'] not in ('cover','close') else str(line)+marker);par.font.name=profile['font_body'];par.font.size=Pt((22 if j==0 else 15) if s['layout']=='statement' else profile['body_pt'] if j else profile['body_pt']+1);par.font.color.rgb=rgb(bodycolor);par.space_after=Pt(12);par.level=0
                if s['layout']=='statement' and j==0: par.font.bold=True;par.space_after=Pt(18)
                if profile.get('rtl') or deck.get('language')=='ar':
                    ppr=par._p.get_or_add_pPr();ppr.set('rtl','1');par.alignment=PP_ALIGN.RIGHT
        if s['layout']=='visual' and s.get('media'):
            img=s['media'][0];candidate=p/img
            if candidate.exists():
                try:
                    from PIL import Image
                    with Image.open(candidate) as im: ratio=im.width/max(im.height,1)
                    maxw=W*.56;maxh=H*.60
                    picw=min(maxw,maxh*ratio);pich=min(maxh,maxw/max(ratio,.01))
                    sl.shapes.add_picture(str(candidate), Inches(W*.39+(maxw-picw)/2), Inches(2.0+(maxh-pich)/2), width=Inches(picw), height=Inches(pich))
                except Exception:
                    try: sl.shapes.add_picture(str(candidate), Inches(W*.39), Inches(2.15), width=Inches(W*.56))
                    except Exception: pass
        if s.get('sources'):
            sb=sl.shapes.add_textbox(x,Inches(H-.94),Inches(W-2*m),Inches(.55));st=sb.text_frame;st.clear();st.word_wrap=True
            rows=[]
            for j,source in enumerate(unique_sources[:8],1):
                where=source.split(' · ',1)[1] if ' · ' in source else ''
                where=where.replace('page ','p.').replace('slide ','s.').replace('paragraph ','para.').replace('row ','row ')
                rows.append(f'[{j}] {source_name(source)} {where}'.strip())
            q=st.paragraphs[0];q.text='Sources: '+'; '.join(rows);q.font.name=profile['font_body'];q.font.size=Pt(8.5);q.font.color.rgb=rgb(c['paper'] if s['layout']=='statement' else c['muted'])
        foot=sl.shapes.add_textbox(x,Inches(H-.38),Inches(W-2*m),Inches(.25));q=foot.text_frame.paragraphs[0];q.text=(profile.get('footer','UHE • DRAFT')+'  |  '+str(i) if profile.get('page_numbers',True) else profile.get('footer','UHE • DRAFT'));q.font.name=profile['font_body'];q.font.size=Pt(8);q.font.color.rgb=rgb(c['paper'] if s['layout']=='statement' else c['muted'])
        sl.notes_slide.notes_text_frame.text='Status: DRAFT • NOT APPROVED\nSources: '+'; '.join(s.get('sources',[]))
    out=p/'output';out.mkdir(exist_ok=True);base=slug(deck['project_name'])+'_DRAFT';idx=1
    while True:
        suffix='' if idx==1 else f'_R{idx:02d}'
        ppt=out/(base+suffix+'.pptx')
        if not ppt.exists(): break
        idx+=1
    r.save(ppt);return ppt

def export_pdf(ppt):
    soffice=shutil.which('soffice')
    if not soffice: return None,'LibreOffice (soffice) not found; PDF export unavailable.'
    out=ppt.parent
    r=subprocess.run([soffice,'-env:UserInstallation=file://'+str(out/'lo_profile'),'--headless','--convert-to','pdf','--outdir',str(out),str(ppt)],capture_output=True,text=True,timeout=120)
    pdf=ppt.with_suffix('.pdf')
    if r.returncode or not pdf.exists(): return None,(r.stderr or r.stdout or 'PDF conversion failed')[-1200:]
    return pdf,None

class Handler(BaseHTTPRequestHandler):
    server_version='UHEStudio/1.0'
    def log_message(self,*a): pass
    def send_json(self,obj,status=200):
        b=json.dumps(obj,ensure_ascii=False).encode();self.send_response(status);self.send_header('Content-Type','application/json; charset=utf-8');self.send_header('Content-Length',str(len(b)));self.end_headers();self.wfile.write(b)
    def do_GET(self):
        u=urlparse(self.path);path=unquote(u.path)
        if path=='/': return self.serve(WEB/'index.html','text/html; charset=utf-8')
        if path.startswith('/web/'):
            asset=(WEB/path.removeprefix('/web/')).resolve()
            if WEB.resolve() not in asset.parents: return self.send_json({'error':'Not found'},404)
            return self.serve(asset,mimetypes.guess_type(path)[0] or 'application/octet-stream')
        if path=='/api/projects':
            return self.send_json([{'project_id':x.name,'name':read_json(x/'inventory.json',{}).get('name',x.name),'files':len(read_json(x/'inventory.json',{}).get('files',[])),'status':'draft' if (x/'slides.json').exists() else 'intake'} for x in sorted(PROJECTS.iterdir()) if x.is_dir()])
        if path=='/api/project':
            pid=parse_qs(u.query).get('id',[''])[0];p=safe_project(pid)
            inv=read_json(p/'inventory.json',None)
            if inv is None: inv=inventory_project(p)
            return self.send_json({'inventory':inv,'deck':read_json(p/'slides.json',None),'profile':read_json(PROFILES/'uhe-default.json',DEFAULT_PROFILE)})
        if path=='/api/profile':
            name=parse_qs(u.query).get('name',['UHE Default'])[0]
            return self.send_json(read_json(PROFILES/(slug(name)+'.json'),read_json(PROFILES/'uhe-default.json',DEFAULT_PROFILE)))
        if path=='/api/profiles':
            return self.send_json([read_json(f,DEFAULT_PROFILE) for f in sorted(PROFILES.glob('*.json'))])
        if path.startswith('/download/'): return self.download(safe_project(path.split('/')[2]),'/'.join(path.split('/')[3:]))
        return self.send_json({'error':'Not found'},404)
    def serve(self,p,ctype):
        if not p.exists(): return self.send_json({'error':'Not found'},404)
        b=p.read_bytes();self.send_response(200);self.send_header('Content-Type',ctype);self.send_header('Content-Length',str(len(b)));self.end_headers();self.wfile.write(b)
    def download(self,p,rel):
        f=(p/rel).resolve()
        if p.resolve() not in f.parents or not f.is_file(): return self.send_json({'error':'Not found'},404)
        return self.serve(f,mimetypes.guess_type(f.name)[0] or 'application/octet-stream')
    def do_POST(self):
        path=urlparse(self.path).path
        try:
            if path=='/api/project':
                n=int(self.headers.get('Content-Length','0'))
                if n>MAX_UPLOAD: return self.send_json({'error':'Upload exceeds 200 MB.'},413)
                data=json.loads(self.rfile.read(n));p=safe_project(data.get('name','project'));p.mkdir(parents=True,exist_ok=True)
                for item in data.get('files',[]):
                    rel=item.get('relativePath') or item.get('name','upload')
                    # Folder uploads retain their relative hierarchy while blocking traversal.
                    rel=rel.replace('\\','/')
                    parts=[x for x in Path(rel).parts if x not in ('..','.') and x not in ('/','\\')]
                    name=Path(item.get('name','upload')).name;ext=Path(name).suffix.lower()
                    if ext not in ALLOWED: continue
                    raw=__import__('base64').b64decode(item.get('data',''),validate=True)
                    dest=p/'sources'/Path(*parts) if parts else p/'sources'/name
                    dest=dest.resolve()
                    if (p/'sources').resolve() not in dest.parents: continue
                    dest.parent.mkdir(parents=True,exist_ok=True)
                    if dest.exists(): dest=dest.with_name(Path(name).stem+'-'+hashlib.sha1(raw).hexdigest()[:6]+ext)
                    dest.write_bytes(raw)
                inv=inventory_project(p);write_json(p/'inventory.json',inv)
                return self.send_json({'inventory':inv,'project_id':p.name})
            n=int(self.headers.get('Content-Length','0'))
            if n>MAX_UPLOAD: return self.send_json({'error':'Request exceeds 200 MB.'},413)
            data=json.loads(self.rfile.read(n));pid=data.get('project_id','');p=safe_project(pid)
            if path=='/api/generate':
                inv=read_json(p/'inventory.json',None)
                if inv is None: inv=inventory_project(p)
                deck=make_draft(inv,data)
                archive_slides_if_changed(p,deck)
                write_json(p/'slides.json',deck);return self.send_json({'deck':deck})
            if path=='/api/save':
                deck=data.get('deck');
                if not isinstance(deck,dict) or not isinstance(deck.get('slides'),list) or len(deck['slides'])>50: return self.send_json({'error':'Invalid slide data or over 50 slides.'},400)
                archive_slides_if_changed(p,deck);write_json(p/'slides.json',deck);return self.send_json({'ok':True})
            if path=='/api/profile':
                prof=data.get('profile',{});label=slug(data.get('profile_name','UHE Default'))
                write_json(PROFILES/(label+'.json'),{**DEFAULT_PROFILE,**prof,'name':data.get('profile_name','UHE Default')});return self.send_json({'ok':True})
            if path in ('/api/export','/api/preview'):
                deck=read_json(p/'slides.json',None)
                if not deck:return self.send_json({'error':'Generate or save a presentation first.'},400)
                if path=='/api/export' and deck.get('review_items') and data.get('review_ack') is not True:return self.send_json({'error':'Review unresolved items and explicitly acknowledge them before export.'},409)
                prof_name=slug(deck.get('profile_name','UHE Default'))+'.json'
                prof=read_json(PROFILES/prof_name,read_json(PROFILES/'uhe-default.json',DEFAULT_PROFILE));ppt=build_pptx(p,deck,prof);pdf,err=export_pdf(ppt)
                return self.send_json({'pptx':f'/download/{p.name}/output/{ppt.name}','pdf':f'/download/{p.name}/output/{pdf.name}' if pdf else None,'pdf_error':err,'slides':len(deck['slides']),'status':deck['status']})
            return self.send_json({'error':'Not found'},404)
        except Exception as e: return self.send_json({'error':f'{type(e).__name__}: {e}'},500)

if __name__=='__main__':
    host=os.environ.get('UHE_STUDIO_HOST','127.0.0.1');port=int(os.environ.get('UHE_STUDIO_PORT','8877'))
    print(f'UHE Presentation Studio local-only server: http://{host}:{port}');print(f'Project data: {DATA}')
    ThreadingHTTPServer((host,port),Handler).serve_forever()
